import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

from lie_detector.audio.transcribe import ts

TOPIC_LABELS = {"intro": "Знакомство", "alcohol": "Алкоголь", "conflict": "Конфликт", "deception": "Обман", "outro": "Завершение"}
ROLE_LABELS = {"I": "Психолог", "R": "Я"}
SHEET, DONE_ROW, HEADER_ROW, DONE_CELL = "Разметка", 5, 6, "F5"
BANNER = [("Выделите те реплики, в которых вы лгали. Заполняйте исключительно жёлтые строки, так как они содержат ваши реплики.", True),
          ("В колонку «Ложь?» впишите Л при лжи, Ч при частичной лжи или оставьте пустой, если это правда.", True),
          ("Напишите по желанию, что было неправдой или «не моя фраза», если перепутаны говорящие.", False)]
HEADER = ["id", "Время", "Кто", "Тема", "Реплика", "Ложь?\n(Л / Ч / пусто)", "Комментарий (необязательно)"]
LIE = {"л", "ложь", "лож", "соврал", "соврала", "1", "+", "x", "х", "да", "l", "lie"}
PARTIAL = {"ч", "частично", "част", "0.5", "0,5", "~", "±", "p", "partial"}
TRUTH = {"п", "правда", "0", "-", "нет", "t", "truth"}
fill = lambda c: PatternFill("solid", fgColor=c)


def normalize_answer(value):
    v = str(value or "").strip().lower().rstrip(".")
    if v in LIE or v.startswith(("лож", "совр")):
        return "Ложь"
    if v in PARTIAL or v.startswith("част"):
        return "Частично"
    if v in TRUTH or v.startswith("правд"):
        return "Правда"
    return None


def write_sheet(dialog, path, title, previous=None):
    previous, wb = previous or {}, Workbook()
    ws = wb.active
    ws.title = SHEET
    for i, (text, bold, color) in enumerate([(title, True, "000000")] + [(t, b, "C62828" if b else "333333") for t, b in BANNER], 1):
        c = ws.cell(i, 2, text)
        c.font, c.alignment = Font(bold=bold, size=13 if i == 1 else 12, color=color), Alignment(wrap_text=True, vertical="center")
        ws.merge_cells(start_row=i, start_column=2, end_row=i, end_column=7)
    ws.cell(DONE_ROW, 2, "Завершение: напишите «да» после того, как вы изучите все свои строки →").font = Font(bold=True, size=12)
    ws.cell(DONE_ROW, 2).alignment = Alignment(horizontal="right")
    ws.merge_cells(start_row=DONE_ROW, start_column=2, end_row=DONE_ROW, end_column=5)
    done = ws[DONE_CELL]
    done.fill, done.border, done.font = fill("C8E6C9"), Border(*(Side(style="medium"),) * 4), Font(bold=True, size=12)
    for col, h in enumerate(HEADER, 1):
        c = ws.cell(HEADER_ROW, col, h)
        c.font, c.fill, c.alignment = Font(bold=True), fill("E0E0E0"), Alignment(wrap_text=True, vertical="center")
    dv = DataValidation(type="list", formula1='"Л,Ч"', allow_blank=True, showInputMessage=True, promptTitle="Ложь?",
                        prompt="Л: соврал(а), Ч: частично, пусто: правда", showErrorMessage=True, errorStyle="warning",
                        error="Впишите Л, Ч или оставьте пустым")
    dv_done = DataValidation(type="list", formula1='"да"', allow_blank=True)
    ws.add_data_validation(dv)
    ws.add_data_validation(dv_done)
    dv_done.add(done)
    thin = Side(style="thin", color="BDBDBD")
    for r, u in enumerate(dialog, HEADER_ROW + 1):
        ans, comment = previous.get(u["id"], (None, None))
        mine = u["role"] == "R"
        values = [u["id"], ts(u["start"])[:8], ROLE_LABELS.get(u["role"], u["role"]), TOPIC_LABELS.get(u["topic"], u["topic"]),
                  u["text"].strip(), {"Ложь": "Л", "Частично": "Ч"}.get(ans) if mine else None, comment]
        for col, v in enumerate(values, 1):
            c = ws.cell(r, col, v)
            c.alignment = Alignment(wrap_text=col in (5, 7), vertical="top")
            if mine:
                c.fill = fill("FFE082" if col == 6 else "FFF4CC")
            else:
                c.font = Font(color="757575")
        if mine:
            a = ws.cell(r, 6)
            a.border, a.font, a.alignment = Border(thin, thin, thin, thin), Font(bold=True, size=12, color="C62828"), Alignment(horizontal="center")
            dv.add(a)
    for col, width in zip("ABCDEFG", [5, 9, 10, 12, 80, 16, 40]):
        ws.column_dimensions[col].width = width
    ws.column_dimensions["A"].hidden = True
    ws.freeze_panes = ws.cell(HEADER_ROW + 1, 1)
    wb.save(path)


def read_sheet(path, assume_done=False):
    ws = load_workbook(path, data_only=True)[SHEET]
    rows = list(ws.iter_rows(values_only=True))
    h = next(i for i, row in enumerate(rows) if row and row[0] == "id")
    done = assume_done or str(ws[DONE_CELL].value or "").strip().lower() in {"да", "yes", "+", "готово", "1"}
    out = []
    for row in rows[h + 1:]:
        if row[0] is None:
            continue
        mine, raw = row[2] == "Я", row[5]
        ans = normalize_answer(raw) if mine else None
        if mine and ans is None and done and not str(raw or "").strip():
            ans = "Правда"
        out.append({"id": int(row[0]), "self_report": ans, "comment": row[6], "raw": raw, "text": row[4], "is_respondent": mine})
    df = pd.DataFrame(out)
    df.attrs["done"] = done
    return df
